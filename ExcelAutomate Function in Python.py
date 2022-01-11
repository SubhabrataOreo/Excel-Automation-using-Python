#!/usr/bin/env python
# coding: utf-8

# In[1]:


#importing all the relevant libraries
import numpy as np
import pandas as pd
import os
import glob
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

#getting the path of the working directory
path = os.getcwd()

#creating a list of all the excel files of the present directory
files = glob.glob(os.path.join(path, "*.xlsx"))

#creating the function to read the excel files, create pivot tables, excel charts and saving those as new excel sheets
def ExcelAutomate(files):
    count=1   #setting a temporary variable (useful for naming the new excel files)
    
    # loop over the list of excel files
    for f in files:
        # read the excel file
        df = pd.read_excel(f)
        #create the pivot table as per instruction
        pivot = df.pivot_table(index =['Product line'], columns=['Gender'],
                           values =['Total'], aggfunc ='sum')
        #getting the path of working directory without the file name
        path='\\'.join(f.split('\\')[0:-1])
        #saving the pivot table as a new excel file
        pivot.to_excel(f"{path}/f{count}.xlsx")
        #loading the new excel workbook
        wb= load_workbook(f"{path}/f{count}.xlsx")
        #gettin access to the active sheet of the workbook
        sh= wb.active
        #creating an object of BarChart module imported previously
        chart= BarChart()
        #creating the excel chart
        labels=Reference(sh, min_col=1, min_row=4, max_row=13)
        values=Reference(sh, min_col=1, min_row=2, max_row=13, max_col=3)
        chart.add_data(values,titles_from_data=True)
        chart.title= "Bar Chart"
        chart.set_categories(labels)
        chart.x_axis.title="Product Line"
        chart.y_axis.title="Total money spent"
        sh.add_chart(chart, "E12")
        #saving the pivot table along with the excel chart with the same name of the new workbook
        wb.save(f"{path}/f{count}.xlsx")
        #increasing the value of the temporary variable
        count+=1
    
#Calling the function
ExcelAutomate(files)

